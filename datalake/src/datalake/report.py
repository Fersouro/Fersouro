"""Camada de relatorios: pastas de trabalho Excel definidas em YAML.

O `export` grava um modelo gold por arquivo, cru -- serve para quem vai levar o
dado para outro lugar. Um *relatorio* e outra coisa: um arquivo com varias abas
(resumo e detalhe), numero formatado como numero se le (R$, %, dd/mm/aaaa),
linha de total e destaque no que precisa de atencao.

Cada arquivo em ``conf/reports/*.yml`` vira um ``.xlsx``:

    name: margem_pecas
    title: Margem de Pecas
    description: Venda, custo e lucro por filial, liquidos de devolucao.
    formats:
      lucro_porcentagem: percentual
    sheets:
      - name: Resumo
        sql: SELECT filial, sum(lucro) AS lucro FROM margem_pecas GROUP BY 1
        totals: [lucro]
        highlights:
          - when: lucro < 0
            style: vermelho
      - name: Detalhe
        model: margem_pecas          # atalho para SELECT * FROM <modelo gold>

O SQL enxerga as mesmas views da gold (silver ``<fonte>__<tabela>`` e os modelos
gold pelo nome), entao um relatorio e SQL comum -- nao ha linguagem nova.

A primeira aba e sempre a Capa: titulo, quando foi gerado e o que tem em cada
aba. Quem recebe a planilha por e-mail costuma nao saber nem de onde ela veio.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

import duckdb
import yaml

from .config import ConfigError, Settings
from .duck import connect, quote, quote_literal
from .export import EXCEL_MAX_ROWS, _cell, _column_widths, gold_models
from .layers.gold import register_silver_views
from .logging_conf import get_logger
from .storage import paths

log = get_logger(__name__)

# Formatos nomeados -> mascara do Excel. A mascara e neutra de idioma: o Excel
# do usuario mostra com a virgula decimal do Windows dele.
FORMATS: dict[str, str] = {
    "texto": "@",
    "inteiro": "#,##0",
    "numero": "#,##0.00",
    "moeda": "R$ #,##0.00",
    # A gold entrega porcentagem ja multiplicada por 100 (12,5 = 12,5%). O '%'
    # entre aspas e literal: se usasse o formato '0.0%' o Excel multiplicaria
    # de novo e mostraria 1250%.
    "percentual": '#,##0.00"%"',
    # Para quando o valor vem entre 0 e 1.
    "fracao": "0.0%",
    "data": "dd/mm/yyyy",
    "data_hora": "dd/mm/yyyy hh:mm",
    "mes": "mm/yyyy",
}

# Estilos de destaque (fundo, fonte) -- as mesmas cores do "Ruim/Bom" do Excel.
STYLES: dict[str, tuple[str, str]] = {
    "vermelho": ("FFC7CE", "9C0006"),
    "amarelo": ("FFEB9C", "9C6500"),
    "verde": ("C6EFCE", "006100"),
    "azul": ("DDEBF7", "1F4E78"),
}

_MONEY_HINTS = (
    "valor", "vlr_", "val_", "preco", "venda", "custo", "lucro", "desconto",
    "frete", "ticket", "faturamento", "receita", "limite_credito",
)
_PERCENT_HINTS = ("porcentagem", "percentual", "_pct", "pct_")
_INT_HINTS = ("qtd", "quantidade", "numero", "nro", "codigo", "id_", "_id")


def _inicio_do_mes() -> dt.date:
    hoje = dt.date.today()
    return dt.date(hoje.year, hoje.month, 1)


def _fim_do_mes() -> dt.date:
    hoje = dt.date.today()
    if hoje.month == 12:
        return dt.date(hoje.year, 12, 31)
    return dt.date(hoje.year, hoje.month + 1, 1) - dt.timedelta(days=1)


# Aceitos no 'default' (e digitaveis no campo) de um parametro do tipo 'data'.
_DATA_ATALHOS = {
    "hoje": dt.date.today,
    "inicio-do-mes": _inicio_do_mes,
    "fim-do-mes": _fim_do_mes,
    "ontem": lambda: dt.date.today() - dt.timedelta(days=1),
}


@dataclass(frozen=True)
class Parameter:
    """Um campo que quem gera o relatorio escolhe (mes, filial, departamento).

    O valor vai para o SQL como parametro nomeado do DuckDB ($nome), nunca por
    concatenacao: assim o formulario da pagina nao consegue injetar SQL.
    """

    name: str
    label: str
    type: str = "texto"          # mes | data | numero | texto | lista
    default: Any = None
    optional: bool = False
    options: tuple[tuple[str, str], ...] = ()   # (valor, rotulo) do tipo 'lista'

    def resolve(self, bruto: Any) -> Any:
        """Converte o que veio do formulario (ou o default) para o tipo certo.

        Nao informado (None) e informado em branco sao coisas diferentes: quem
        nao mandou o campo quer o default; quem apagou o 410 do formulario quer
        'todos os departamentos', e cair no default de novo ignoraria o pedido.
        """
        if bruto is None:
            bruto = self.default
        vazio = bruto is None or (isinstance(bruto, str) and not bruto.strip())
        if vazio:
            if self.optional:
                return None
            raise ValueError(f"parametro '{self.name}' ({self.label}) e obrigatorio")

        texto = str(bruto).strip()
        if self.type == "lista":
            validos = {v for v, _ in self.options}
            if texto not in validos:
                rotulos = ", ".join(f"{v} ({r})" for v, r in self.options)
                raise ValueError(
                    f"parametro '{self.name}': '{texto}' nao e uma opcao valida. "
                    f"Escolha: {rotulos}"
                )
            try:                       # numero quando o valor e numero
                return int(texto)
            except ValueError:
                return texto
        if self.type == "mes":
            # 'atual' e o caso comum: o relatorio do mes corrente, sem digitar.
            if texto.lower() in ("atual", "corrente", "hoje"):
                hoje = dt.date.today()
                return dt.date(hoje.year, hoje.month, 1)
            for formato in ("%Y-%m", "%m/%Y", "%Y-%m-%d"):
                try:
                    achado = dt.datetime.strptime(texto, formato).date()
                    return dt.date(achado.year, achado.month, 1)
                except ValueError:
                    continue
            raise ValueError(
                f"parametro '{self.name}': '{texto}' nao e uma competencia "
                f"(use AAAA-MM, por exemplo 2026-09)"
            )
        if self.type == "data":
            # Atalhos para o default: um relatorio de periodo quase sempre abre
            # no mes corrente, e escrever a data na mao a cada geracao cansa.
            atalho = _DATA_ATALHOS.get(texto.lower())
            if atalho:
                return atalho()
            for formato in ("%Y-%m-%d", "%d/%m/%Y"):
                try:
                    return dt.datetime.strptime(texto, formato).date()
                except ValueError:
                    continue
            raise ValueError(
                f"parametro '{self.name}': '{texto}' nao e uma data "
                f"(use AAAA-MM-DD, por exemplo 2026-09-30)"
            )
        if self.type == "numero":
            try:
                return int(texto)
            except ValueError:
                try:
                    return float(texto)
                except ValueError:
                    raise ValueError(
                        f"parametro '{self.name}': '{texto}' nao e um numero"
                    ) from None
        return texto


@dataclass(frozen=True)
class Highlight:
    """Regra de destaque: uma condicao SQL e a cor que ela pinta."""

    when: str
    style: str = "vermelho"
    scope: str = "row"          # 'row' ou o nome de uma coluna


@dataclass(frozen=True)
class SheetConfig:
    name: str
    sql: str
    description: str | None = None
    totals: tuple[str, ...] = ()
    highlights: tuple[Highlight, ...] = ()
    formats: dict[str, str] = field(default_factory=dict)
    limit: int | None = None


@dataclass(frozen=True)
class ReportConfig:
    name: str
    title: str
    description: str | None = None
    sheets: tuple[SheetConfig, ...] = ()
    formats: dict[str, str] = field(default_factory=dict)
    parameters: tuple[Parameter, ...] = ()
    path: Path | None = None

    def resolve_parameters(self, valores: dict[str, Any] | None = None) -> dict[str, Any]:
        """Valores prontos para o DuckDB, com os defaults preenchidos."""
        valores = valores or {}
        desconhecidos = set(valores) - {p.name for p in self.parameters}
        if desconhecidos:
            raise ValueError(
                f"[{self.name}] parametro inexistente: {', '.join(sorted(desconhecidos))}. "
                f"Existem: {', '.join(p.name for p in self.parameters) or '(nenhum)'}"
            )
        return {p.name: p.resolve(valores.get(p.name)) for p in self.parameters}

    @classmethod
    def from_dict(cls, data: dict[str, Any], path: Path | None = None) -> "ReportConfig":
        name = str(data.get("name") or (path.stem if path else "")).strip().lower()
        if not name:
            raise ConfigError(f"Relatorio sem 'name' em {path}")

        raw_sheets = data.get("sheets") or []
        if not raw_sheets:
            raise ConfigError(f"Relatorio '{name}' sem 'sheets'")

        formats = {k.lower(): str(v) for k, v in (data.get("formats") or {}).items()}

        parametros: list[Parameter] = []
        for raw in data.get("parameters") or []:
            if not isinstance(raw, dict) or not raw.get("name"):
                raise ConfigError(f"[{name}] parametro sem 'name'")
            tipo = str(raw.get("type") or "texto").lower()
            if tipo not in ("mes", "data", "numero", "texto", "lista"):
                raise ConfigError(
                    f"[{name}.{raw['name']}] tipo '{tipo}' invalido; "
                    f"use mes, data, numero, texto ou lista"
                )
            opcoes: list[tuple[str, str]] = []
            for op in raw.get("options") or []:
                if isinstance(op, dict):
                    valor = "" if op.get("value") is None else str(op["value"])
                    opcoes.append((valor, str(op.get("label") or valor)))
                else:
                    opcoes.append((str(op), str(op)))
            if tipo == "lista" and not opcoes:
                raise ConfigError(f"[{name}.{raw['name']}] tipo 'lista' exige 'options'")
            parametros.append(
                Parameter(
                    name=str(raw["name"]).strip(),
                    label=str(raw.get("label") or raw["name"]),
                    type=tipo,
                    default=raw.get("default"),
                    optional=bool(raw.get("optional")),
                    options=tuple(opcoes),
                )
            )
        sheets: list[SheetConfig] = []
        for i, raw in enumerate(raw_sheets, start=1):
            if not isinstance(raw, dict):
                raise ConfigError(f"[{name}] aba {i}: esperado um mapeamento YAML")
            aba = str(raw.get("name") or f"Aba {i}").strip()
            sql = (raw.get("sql") or "").strip().rstrip(";")
            model = (raw.get("model") or "").strip()
            if sql and model:
                raise ConfigError(f"[{name}.{aba}] use 'sql' ou 'model', nao os dois")
            if not sql:
                if not model:
                    raise ConfigError(f"[{name}.{aba}] falta 'sql' ou 'model'")
                sql = f"SELECT * FROM {quote(model.lower())}"

            highlights = []
            for regra in raw.get("highlights") or []:
                if not isinstance(regra, dict) or not regra.get("when"):
                    raise ConfigError(f"[{name}.{aba}] destaque sem 'when'")
                estilo = str(regra.get("style") or "vermelho").lower()
                if estilo not in STYLES:
                    raise ConfigError(
                        f"[{name}.{aba}] estilo '{estilo}' invalido; "
                        f"use um de {', '.join(sorted(STYLES))}"
                    )
                highlights.append(
                    Highlight(
                        when=str(regra["when"]).strip().rstrip(";"),
                        style=estilo,
                        scope=str(regra.get("scope") or "row"),
                    )
                )

            limite = raw.get("limit")
            sheets.append(
                SheetConfig(
                    name=aba,
                    sql=sql,
                    description=raw.get("description"),
                    totals=tuple(str(c).lower() for c in (raw.get("totals") or ())),
                    highlights=tuple(highlights),
                    formats={
                        k.lower(): str(v) for k, v in (raw.get("formats") or {}).items()
                    },
                    limit=int(limite) if limite else None,
                )
            )

        vistas: set[str] = set()
        for sheet in sheets:
            chave = sheet.name.lower()[:31]
            if chave in vistas:
                raise ConfigError(f"[{name}] duas abas com o nome '{sheet.name}'")
            vistas.add(chave)

        return cls(
            name=name,
            title=str(data.get("title") or name.replace("_", " ").title()),
            description=data.get("description"),
            sheets=tuple(sheets),
            formats=formats,
            parameters=tuple(parametros),
            path=path,
        )


@dataclass
class ReportResult:
    report: str
    path: Path | None
    rows: int
    status: str
    message: str | None = None

    @property
    def ok(self) -> bool:
        return self.status in ("success", "skipped")


def reports_dir(settings: Settings) -> Path:
    return settings.project_root / "conf" / "reports"


def load_reports(settings: Settings) -> list[ReportConfig]:
    """Le todos os conf/reports/*.yml, em ordem de nome de arquivo."""
    directory = reports_dir(settings)
    if not directory.is_dir():
        return []
    relatorios: list[ReportConfig] = []
    for path in sorted(directory.glob("*.y*ml")):
        if path.name.startswith("_"):
            continue
        with path.open("r", encoding="utf-8") as handle:
            data = yaml.safe_load(handle) or {}
        if not isinstance(data, dict):
            raise ConfigError(f"Esperado um mapeamento YAML em {path}")
        relatorios.append(ReportConfig.from_dict(data, path))
    nomes: set[str] = set()
    for relatorio in relatorios:
        if relatorio.name in nomes:
            raise ConfigError(f"Relatorio '{relatorio.name}' definido duas vezes")
        nomes.add(relatorio.name)
    return relatorios


def register_history_views(con, settings: Settings) -> list[str]:
    """Registra o historico diario de estoque, se existir.

    Os snapshots que a pagina de Estoque Minimo grava
    (<lake>/historico_estoque/AAAA-MM-DD.parquet, com a coluna 'data') sao a
    unica serie temporal de saldo que existe: o ERP so guarda o saldo de agora.
    Registrando aqui, um relatorio pode perguntar "como estava em tal dia".
    """
    pasta = settings.root / "historico_estoque"
    if not pasta.is_dir() or not any(pasta.glob("*.parquet")):
        return []
    relation = f"read_parquet({quote_literal(str(pasta / '*.parquet'))})"
    con.execute(
        f"CREATE OR REPLACE VIEW {quote('historico_estoque')} AS SELECT * FROM {relation}"
    )
    return ["historico_estoque"]


def register_gold_views(con, settings: Settings) -> list[str]:
    """Cria uma view por modelo gold materializado."""
    nomes: list[str] = []
    for nome, directory in gold_models(settings).items():
        relation = f"read_parquet({quote_literal(paths.glob_parquet(directory))})"
        con.execute(f"CREATE OR REPLACE VIEW {quote(nome)} AS SELECT * FROM {relation}")
        nomes.append(nome)
    return nomes


# ---------------------------------------------------------------- formatacao


def _resolve_format(nome: str) -> str:
    """Aceita um formato nomeado ('moeda') ou uma mascara do Excel crua."""
    chave = nome.strip().lower()
    if chave in FORMATS:
        return FORMATS[chave]
    return nome         # mascara literal, ex.: '#,##0.000'


def infer_format(coluna: str, amostra: list[Any]) -> str | None:
    """Formato provavel de uma coluna, pelo nome e pelos valores.

    O nome vem primeiro porque um float chamado 'venda_total' e dinheiro, e um
    float chamado 'lucro_porcentagem' e porcentagem -- o tipo nao distingue.
    """
    nome = coluna.lower()
    valor = next((v for v in amostra if v is not None), None)

    if isinstance(valor, bool) or valor is None:
        return None

    if isinstance(valor, (dt.date, dt.datetime)):
        if nome.startswith("competencia") or nome in ("mes", "mes_ano"):
            return FORMATS["mes"]
        if isinstance(valor, dt.datetime):
            return FORMATS["data_hora"]
        return FORMATS["data"]

    if not isinstance(valor, (int, float, Decimal)):
        return None

    if any(h in nome for h in _PERCENT_HINTS):
        return FORMATS["percentual"]
    # Inteiro antes de dinheiro de proposito: 'revenda' contem 'venda' mas e o
    # numero da loja, e dinheiro no lake sempre vem como decimal.
    if isinstance(valor, int) or any(
        nome.startswith(h) or nome.endswith(h) for h in _INT_HINTS
    ):
        return FORMATS["inteiro"]
    if any(h in nome for h in _MONEY_HINTS):
        return FORMATS["moeda"]
    return FORMATS["numero"]


def column_formats(
    colunas: list[str], linhas: list[tuple], escolhidos: dict[str, str]
) -> list[str | None]:
    """Formato por coluna: o declarado no YAML vence; o resto e inferido."""
    amostra = linhas[:200]
    formatos: list[str | None] = []
    for i, coluna in enumerate(colunas):
        declarado = escolhidos.get(coluna.lower())
        if declarado:
            formatos.append(_resolve_format(declarado))
        else:
            formatos.append(infer_format(coluna, [linha[i] for linha in amostra]))
    return formatos


# ------------------------------------------------------------------- leitura


def _sheet_sql(sheet: SheetConfig, limite: int) -> str:
    """SQL da aba com as condicoes de destaque como colunas extras.

    Avaliar a condicao no DuckDB (e nao em Python) e o que permite escrever
    qualquer expressao SQL no 'when' sem inventar um mini-interpretador.
    """
    if not sheet.highlights:
        base = f"SELECT * FROM ({sheet.sql}) AS _aba"
    else:
        extras = ", ".join(
            f"CAST(({h.when}) AS BOOLEAN) AS {quote(f'__hl_{i}')}"
            for i, h in enumerate(sheet.highlights)
        )
        base = f"SELECT _aba.*, {extras} FROM ({sheet.sql}) AS _aba"
    return f"{base} LIMIT {int(limite)}"


def fetch_sheet(con, sheet: SheetConfig, max_rows: int, parametros=None):
    """Executa a aba. -> (colunas, linhas, marcas, total_real).

    O corte vai no SQL, nao em Python: uma aba que consultasse dez milhoes de
    linhas para jogar fora nove nao caberia na memoria.
    """
    limite = min(sheet.limit, max_rows) if sheet.limit else max_rows
    sql = _sheet_sql(sheet, limite)
    resultado = con.execute(sql, parametros) if parametros else con.execute(sql)
    colunas_todas = [d[0] for d in resultado.description]
    dados = resultado.fetchall()

    total = len(dados)
    if len(dados) == max_rows and (sheet.limit is None or sheet.limit > max_rows):
        # Bateu no teto do formato -- so aqui vale pagar uma contagem para
        # dizer quantas linhas ficaram de fora.
        contagem = f"SELECT count(*) FROM ({sheet.sql}) AS _aba"
        total = (
            con.execute(contagem, parametros) if parametros else con.execute(contagem)
        ).fetchone()[0]

    n_hl = len(sheet.highlights)
    colunas = colunas_todas[: len(colunas_todas) - n_hl] if n_hl else colunas_todas
    if n_hl:
        marcas = [[bool(linha[-n_hl + i]) for linha in dados] for i in range(n_hl)]
        linhas = [linha[:-n_hl] for linha in dados]
    else:
        marcas = []
        linhas = list(dados)

    faltando = [c for c in sheet.totals if c not in {x.lower() for x in colunas}]
    if faltando:
        raise ValueError(
            f"aba '{sheet.name}': coluna de total inexistente: {', '.join(faltando)}"
        )
    return colunas, linhas, marcas, total


# ------------------------------------------------------------------- escrita


def _write_cover(wb, relatorio: ReportConfig, resumo, parametros=None) -> None:
    """Capa: o que e, quando foi gerado e o que tem em cada aba."""
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = wb.create_sheet("Capa", 0)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 14

    ws["A1"] = relatorio.title
    ws["A1"].font = Font(bold=True, size=16)
    linha = 3
    if relatorio.description:
        ws.cell(linha, 1, str(relatorio.description).strip())
        ws.cell(linha, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=3)
        ws.row_dimensions[linha].height = 32
        linha += 2

    ws.cell(linha, 1, "Gerado em").font = Font(bold=True)
    ws.cell(linha, 2, dt.datetime.now().strftime("%d/%m/%Y %H:%M"))
    linha += 1
    ws.cell(linha, 1, "Origem").font = Font(bold=True)
    ws.cell(linha, 2, "datalake (camada gold)")
    linha += 1

    # Sem isto, duas planilhas do mesmo relatorio com filtros diferentes ficam
    # indistinguiveis depois de salvas.
    for parametro in relatorio.parameters:
        valor = (parametros or {}).get(parametro.name)
        if isinstance(valor, dt.date):
            mostrado = valor.strftime("%m/%Y" if parametro.type == "mes" else "%d/%m/%Y")
        elif parametro.options:
            # Na capa vale o que a pessoa escolheu na tela ("Consolidado"), nao
            # o valor tecnico que foi para o SQL.
            escolhido = "" if valor is None else str(valor)
            mostrado = next((r for v, r in parametro.options if v == escolhido), escolhido or "(todos)")
        else:
            mostrado = "(todos)" if valor is None else str(valor)
        ws.cell(linha, 1, parametro.label).font = Font(bold=True)
        ws.cell(linha, 2, mostrado)
        linha += 1
    linha += 1

    for coluna, titulo in enumerate(("Aba", "O que mostra", "Linhas"), start=1):
        celula = ws.cell(linha, coluna, titulo)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="44546A")
    linha += 1
    for nome, descricao, linhas_aba in resumo:
        ws.cell(linha, 1, nome)
        ws.cell(linha, 2, descricao or "")
        ws.cell(linha, 3, linhas_aba).number_format = FORMATS["inteiro"]
        linha += 1


def _write_sheet(wb, sheet: SheetConfig, colunas, linhas, marcas, formatos) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(sheet.name[:31] or "dados")
    ws.append(list(colunas))
    for linha in linhas:
        ws.append([_cell(v) for v in linha])

    cabecalho_font = Font(bold=True, color="FFFFFF")
    cabecalho_fill = PatternFill("solid", fgColor="44546A")
    for celula in ws[1]:
        celula.font = cabecalho_font
        celula.fill = cabecalho_fill
        celula.alignment = Alignment(vertical="center", wrap_text=True)

    for i, formato in enumerate(formatos, start=1):
        if not formato:
            continue
        letra = get_column_letter(i)
        for celula in ws[letra][1:]:
            celula.number_format = formato

    for i, largura in enumerate(_column_widths(colunas, linhas), start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura

    ws.freeze_panes = "A2"
    ultima = len(linhas) + 1
    if linhas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{ultima}"

    # Destaques: pintam a linha (ou uma celula) conforme a condicao do YAML.
    indice = {c.lower(): i + 1 for i, c in enumerate(colunas)}
    for regra, marcado in zip(sheet.highlights, marcas):
        fundo, cor = STYLES[regra.style]
        fill = PatternFill("solid", fgColor=fundo)
        fonte = Font(color=cor)
        alvo = indice.get(regra.scope.lower()) if regra.scope != "row" else None
        for i, marca in enumerate(marcado):
            if not marca:
                continue
            linha_excel = i + 2
            faixa = [alvo] if alvo else range(1, len(colunas) + 1)
            for coluna in faixa:
                celula = ws.cell(linha_excel, coluna)
                celula.fill = fill
                celula.font = fonte

    if sheet.totals and linhas:
        # SUBTOTAL(109) em vez de SUM: com o filtro ligado, o total passa a ser
        # o do que esta filtrado -- que e o numero que a pessoa esta olhando.
        total_linha = ultima + 1
        ws.cell(total_linha, 1, "TOTAL").font = Font(bold=True)
        for coluna in sheet.totals:
            i = indice[coluna]
            letra = get_column_letter(i)
            celula = ws.cell(total_linha, i, f"=SUBTOTAL(109,{letra}2:{letra}{ultima})")
            celula.font = Font(bold=True)
            celula.number_format = formatos[i - 1] or FORMATS["numero"]


def build_report(
    settings: Settings,
    relatorio: ReportConfig,
    con,
    destino_dir: Path,
    valores: dict[str, Any] | None = None,
) -> ReportResult:
    """Executa todas as abas e grava um .xlsx."""
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "Pacote 'openpyxl' nao instalado. Rode: pip install openpyxl"
        ) from exc

    destino = destino_dir / f"{relatorio.name}.xlsx"
    avisos: list[str] = []
    total_linhas = 0

    try:
        parametros = relatorio.resolve_parameters(valores)
        wb = Workbook()
        wb.remove(wb.active)
        resumo: list[tuple[str, str, int]] = []

        for sheet in relatorio.sheets:
            limite = EXCEL_MAX_ROWS - (1 if sheet.totals else 0)
            colunas, linhas, marcas, total = fetch_sheet(con, sheet, limite, parametros)
            if total > len(linhas):
                aviso = (
                    f"aba '{sheet.name}': {total:,} linhas nao cabem no xlsx; "
                    f"gravadas as primeiras {len(linhas):,}"
                )
                avisos.append(aviso)
                log.warning("[report.%s] %s", relatorio.name, aviso)

            formatos = column_formats(
                colunas, linhas, {**relatorio.formats, **sheet.formats}
            )
            _write_sheet(wb, sheet, colunas, linhas, marcas, formatos)
            resumo.append((sheet.name[:31], sheet.description or "", len(linhas)))
            total_linhas += len(linhas)

        _write_cover(wb, relatorio, resumo, parametros)
        destino.parent.mkdir(parents=True, exist_ok=True)
        wb.save(destino)

        log.info(
            "[report.%s] %s abas, %s linhas -> %s",
            relatorio.name, len(relatorio.sheets), f"{total_linhas:,}", destino,
        )
        return ReportResult(
            relatorio.name, destino, total_linhas, "success", "; ".join(avisos) or None
        )

    except duckdb.CatalogException as exc:
        # O relatorio cita um modelo que ainda nao existe na gold. Igual a gold:
        # e rotina num lake com varias fontes, nao e falha.
        faltante = str(exc).split("\n")[0]
        log.warning("[report.%s] ignorado: %s", relatorio.name, faltante)
        return ReportResult(
            relatorio.name, None, 0, "skipped", "depende de modelo ausente na gold"
        )
    except Exception as exc:  # noqa: BLE001
        log.error("[report.%s] falhou: %s", relatorio.name, exc)
        return ReportResult(
            relatorio.name, None, 0, "failed", f"{type(exc).__name__}: {exc}"
        )


def build_all(
    settings: Settings,
    apenas: list[str] | None = None,
    destino_dir: Path | None = None,
    valores: dict[str, Any] | None = None,
) -> list[ReportResult]:
    relatorios = load_reports(settings)
    if apenas:
        querido = {r.lower() for r in apenas}
        disponiveis = {r.name for r in relatorios}
        desconhecidos = querido - disponiveis
        if desconhecidos:
            raise ValueError(
                f"Relatorio inexistente: {', '.join(sorted(desconhecidos))}. "
                f"Disponiveis: {', '.join(sorted(disponiveis)) or '(nenhum)'}"
            )
        relatorios = [r for r in relatorios if r.name in querido]
    if not relatorios:
        log.warning("Nenhum relatorio em %s", reports_dir(settings))
        return []

    destino = destino_dir or settings.reports_dir
    con = connect(settings)
    try:
        register_silver_views(con, settings)
        register_gold_views(con, settings)
        register_history_views(con, settings)
        return [build_report(settings, r, con, destino, valores) for r in relatorios]
    finally:
        con.close()
