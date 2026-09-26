"""Fluxo completo contra um portal FALSO local (login, menu, nova aba, iframe,
lista, downloads). Exige Playwright + um Chromium; sem eles, e pulado."""

from __future__ import annotations

import http.server
import os
import threading
from decimal import Decimal
from functools import partial
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from pdf_fake import gerar_pdf, pdf_sg

pytest.importorskip("playwright.sync_api")
from datalake.rpa import saga_vh47  # noqa: E402
from datalake.rpa.controle import Controle, ItemLista  # noqa: E402


def _achar_chromium() -> str | None:
    if os.environ.get("RPA_CHROMIUM_EXECUTAVEL"):
        return os.environ["RPA_CHROMIUM_EXECUTAVEL"]
    for p in sorted(Path("/opt/pw-browsers").glob("chromium-*/chrome-linux*/chrome")):
        return str(p)
    return None


CHROMIUM = _achar_chromium()
pytestmark = pytest.mark.skipif(not CHROMIUM, reason="sem Chromium para o Playwright")

INDEX = """<html><head><title>Portal Rede VW</title></head><body>
<script>
if (!document.cookie.includes('logado=1')) {
  document.write('<form onsubmit="document.cookie=\\'logado=1\\'; location.reload(); return false">'
   + '<input name="username" placeholder="Usuario"><input type="password" name="senha">'
   + '<button type="submit">Entrar</button></form>');
} else {
  document.write('<div id="menu"><a href="#">Manual da GARANTIA VOLKSWAGEN</a> '
   + '<a href="#" onclick="document.getElementById(\\'sub\\').style.display=\\'block\\'">GARANTIA VOLKSWAGEN</a>'
   + '<ul id="sub" style="display:none"><li><a href="#">SAGA Manual</a></li>'
   + '<li><a href="/saga.html" target="_blank">SAGA</a></li></ul></div>');
}
</script></body></html>"""

SAGA = """<html><body><h1>SAGA</h1><iframe src="/conteudo.html" width="800" height="600"></iframe></body></html>"""

CONTEUDO = """<html><body><table>
<tr><th>Relatorio</th><th>Arquivos</th></tr>
<tr><td>SAGA1 - VH12</td><td><a href="/errado.html">Lista de arquivos</a></td></tr>
<tr><td>SAGA2 - VH47</td><td><a href="/lista.html">Lista de arquivos</a></td></tr>
</table></body></html>"""

ERRADO = "<html><body>relatorio errado</body></html>"

ARQUIVOS = [  # ano, mes, regional, dn, nome, conteudo
    (2025, "Agosto", "SUL", "1234", "VH47_0825_01.pdf", ("100001", "1.000,00")),
    (2025, "Agosto", "SUL", "1234", "VH47_0825_02.pdf", ("100002", "R$ 2.500,50")),
    (2025, "Setembro", "SUL", "1234", "VH47_0925_01.pdf", ("100003", "300,00")),
    (2025, "Setembro", "SUL", "1234", "VH47_0925_02.pdf", ("100004", "400,00")),
    (2025, "Setembro", "SUL", "1234", "VH47_0925_03.pdf", ("100005", "-50,00")),
    (2025, "Setembro", "SUL", "1234", "VH47_0925_04.pdf", "sem_total"),
    (2025, "Setembro", "SUL", "1234", "VH47_0925_05.pdf", "html"),
    (2025, "Setembro", "NORTE", "9999", "VH47_0925_99.pdf", ("999999", "9,99")),
]


def _lista_html() -> str:
    linhas = "".join(
        f"<tr><td>{a}</td><td>{m}</td><td>{r}</td><td>{dn}</td>"
        f"<td><a href='/pdf/{n}'>{n}</a></td></tr>" for a, m, r, dn, n, _ in reversed(ARQUIVOS)
    )
    return ("<html><body><h2>Lista de arquivos - SAGA2 - VH47</h2><table>"
            "<tr><th>Ano</th><th>Mês</th><th>Regional</th><th>DN</th><th>Nome do arquivo</th></tr>"
            f"{linhas}</table></body></html>")


class _Handler(http.server.SimpleHTTPRequestHandler):
    def end_headers(self):
        if self.path.startswith("/pdf/"):
            self.send_header("Content-Disposition", "attachment")
        super().end_headers()

    def log_message(self, *a):
        pass


@pytest.fixture()
def portal(tmp_path):
    raiz = tmp_path / "site"
    (raiz / "pdf").mkdir(parents=True)
    for nome, html in {"index.html": INDEX, "saga.html": SAGA, "conteudo.html": CONTEUDO,
                       "errado.html": ERRADO, "lista.html": _lista_html()}.items():
        (raiz / nome).write_text(html, encoding="utf-8")
    for *_, nome, conteudo in ARQUIVOS:
        destino = raiz / "pdf" / nome
        if conteudo == "html":
            destino.write_text("<html>sessao expirada</html>")
        elif conteudo == "sem_total":
            gerar_pdf(destino, [[(50, 800, "Nº SG: 100006"), (50, 780, "Peças: 10,00")]])
        else:
            pdf_sg(destino, *conteudo)
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), partial(_Handler, directory=str(raiz)))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_address[1]}/"
    srv.shutdown()


@pytest.fixture()
def ambiente(tmp_path, portal, monkeypatch):
    xl = tmp_path / "garantia.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["O.S.", "NF de mão de obra (serviço)", "NF de peças", "Data",
               "Valor de mão de obra", "Valor de peça", "Valor total da SG"])
    ws.append([100001, "NF1", None, None, None, None, 1000.0])  # ja lancada a mao
    wb.save(xl)
    conf = Path(saga_vh47.CONFIG_PADRAO).read_text(encoding="utf-8")
    conf = conf.replace("https://www.portalredevw.com.br/", portal)
    conf = conf.replace("navegador: msedge", "navegador: chromium").replace("headless: false", "headless: true")
    conf = conf.replace("timeout_s: 60", "timeout_s: 10").replace("timeout_download_s: 120", "timeout_download_s: 15")
    cfg = tmp_path / "saga.yml"
    cfg.write_text(conf, encoding="utf-8")
    monkeypatch.setenv("SAGA_DN", "1234")
    monkeypatch.setenv("SAGA_PLANILHA", str(xl))
    monkeypatch.setenv("RPA_PORTALVW_USUARIO", "teste")
    monkeypatch.setenv("RPA_PORTALVW_SENHA", "teste")
    monkeypatch.setenv("RPA_CHROMIUM_EXECUTAVEL", CHROMIUM)
    args = ["--config", str(cfg), "--pasta", str(tmp_path / "trabalho")]
    return args, xl, tmp_path / "trabalho"


def _sgs(xl: Path) -> list:
    ws = load_workbook(xl).active
    return [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]


def test_fluxo_completo_e_reexecucao(ambiente, caplog):
    args, xl, trabalho = ambiente

    assert saga_vh47.main(args) == 2  # dois arquivos com erro (sem total / nao-PDF)
    assert _sgs(xl) == [100001, 100002, 100003, 100004, 100005]
    ws = load_workbook(xl).active
    assert ws["G3"].value == pytest.approx(2500.50)
    assert ws["G6"].value == pytest.approx(-50.0)
    assert ws["B3"].value is None  # NF nao inventada
    assert ws["B2"].value == "NF1"  # dado manual preservado

    ctl = Controle(trabalho / "controle.sqlite")
    cont = ctl.contagem()
    assert cont == {"processado": 5, "erro": 2}
    with ctl._con() as con:
        erros = {r["nome"]: r["erro"] for r in con.execute("SELECT nome, erro FROM arquivos WHERE status='erro'")}
    assert "Valor total" in erros["VH47_0925_04.pdf"]
    assert "nao e um PDF" in erros["VH47_0925_05.pdf"]
    # o PDF com erro de leitura fica guardado para analise
    assert (trabalho / "pdfs" / "2025" / "09" / "1234" / "VH47_0925_04.pdf").is_file()
    # DN de outra concessionaria nao entra
    assert 999999 not in _sgs(xl)

    # 2a execucao: nada novo -> nada inserido, nenhum download dos ja processados
    caplog.clear()
    assert saga_vh47.main(args) == 2
    assert _sgs(xl) == [100001, 100002, 100003, 100004, 100005]
    texto = caplog.text
    assert "Identificados 2 arquivo(s) ainda nao processados" in texto
    assert "PDF ja baixado antes, reaproveitado: VH47_0925_04.pdf" in texto
    assert "Registro inserido" not in texto


def test_cenario_b_baixa_faltantes_em_ordem(ambiente, caplog):
    args, xl, trabalho = ambiente
    ctl = Controle(trabalho / "controle.sqlite")
    # setembro 01 e 02 ja processados; faltam agosto inteiro e setembro 03..05
    for nome in ("VH47_0925_01.pdf", "VH47_0925_02.pdf"):
        it = ItemLista(2025, 9, "SUL", "1234", nome)
        ctl.registrar_listagem([it])
        ctl.marcar_processado(it.id, "padrao", [])
    saga_vh47.main(args)
    baixados = [l.split("Iniciando download: ")[1] for l in caplog.text.splitlines() if "Iniciando download" in l]
    assert baixados == ["VH47_0825_01.pdf", "VH47_0825_02.pdf", "VH47_0925_03.pdf",
                        "VH47_0925_04.pdf", "VH47_0925_05.pdf"]
    assert 100003 not in _sgs(xl) and 100005 in _sgs(xl)


def test_so_listar_nao_baixa_nem_grava(ambiente, caplog):
    args, xl, trabalho = ambiente
    antes = xl.read_bytes()
    assert saga_vh47.main(args + ["--so-listar"]) == 0
    assert "Identificados 7 arquivo(s)" in caplog.text
    assert xl.read_bytes() == antes
    assert not (trabalho / "pdfs").exists()


def test_importar_pdfs_manuais_idempotente(ambiente, tmp_path):
    args, xl, _ = ambiente
    pasta = tmp_path / "manuais"
    pdf_sg(pasta / "a.pdf", "200001", "1,00")
    pdf_sg(pasta / "b.pdf", "200002", "2,00")
    assert saga_vh47.main(args + ["--importar", str(pasta)]) == 0
    assert saga_vh47.main(args + ["--importar", str(pasta)]) == 0
    assert _sgs(xl) == [100001, 200001, 200002]
    assert load_workbook(xl).active["G4"].value == float(Decimal("2.00"))
