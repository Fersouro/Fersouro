import hashlib
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from datalake.rpa.planilha import ConfigPlanilha, ErroPlanilha, Planilha

CABECALHO = ["O.S.", "NF de mão de obra (serviço)", "NF de peças", "Data",
             "Valor de mão de obra", "Valor de peça", "Valor total da SG", "Conferência"]
COLUNAS = {
    "sg": ["O.S.", "OS"],
    "nf_mao_obra": ["NF de mão de obra (serviço)"],
    "nf_pecas": ["NF de peças"],
    "data": ["Data"],
    "valor_mao_obra": ["Valor de mão de obra"],
    "valor_pecas": ["Valor de peça"],
    "valor_total": ["Valor total da SG"],
}
TIPOS = {"sg": "documento", "valor_total": "dinheiro"}


def criar_planilha(caminho: Path, com_tabela: bool = False) -> Path:
    wb = Workbook()
    capa = wb.active
    capa.title = "Capa"
    capa["A1"] = "Controle de garantia"
    ws = wb.create_sheet("Garantia")
    ws["A1"] = "GRUPO TERRASUL - GARANTIA VW"
    for i, nome in enumerate(CABECALHO, start=1):
        ws.cell(row=3, column=i, value=nome.upper() if i == 1 else nome)
    ws.append([111111, "NF-10", "NF-20", None, 100.0, 200.0, 300.0, "=E4+F4-G4"])
    ws.append([222222, None, None, None, None, None, None, "=E5+F5-G5"])
    ws["G4"].number_format = "#,##0.00"
    ws["A10"].number_format = "@"  # formatacao solta abaixo dos dados nao pode confundir
    if com_tabela:
        ws.add_table(Table(displayName="Garantia", ref="A3:H5"))
    wb.save(caminho)
    return caminho


def cfg(caminho: Path, tmp_path: Path, **extra) -> ConfigPlanilha:
    d = {"caminho": str(caminho), "colunas": COLUNAS, "backup_dir": str(tmp_path / "bkp"), **extra}
    return ConfigPlanilha.from_dict(d, TIPOS, tmp_path)


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def test_insere_sg_nova_como_numero_sem_inventar_campos(tmp_path):
    xl = criar_planilha(tmp_path / "g.xlsx")
    acoes = Planilha(cfg(xl, tmp_path)).upsert([{"sg": "333333", "valor_total": Decimal("1245.80")}])
    assert [(a.chave, a.acao, a.linha) for a in acoes] == [("333333", "inserida", 6)]
    ws = load_workbook(xl)["Garantia"]
    assert ws["A6"].value == 333333
    assert ws["G6"].value == pytest.approx(1245.80) and isinstance(ws["G6"].value, float)
    assert ws["G6"].number_format == "#,##0.00"
    assert all(ws.cell(row=6, column=c).value is None for c in (2, 3, 4, 5, 6))
    assert ws["H6"].value == "=E6+F6-G6"  # formula da linha de cima, ajustada
    # nada do que existia mudou
    assert [c.value for c in ws[4]][:7] == [111111, "NF-10", "NF-20", None, 100, 200, 300]
    assert load_workbook(xl)["Capa"]["A1"].value == "Controle de garantia"


def test_reexecucao_nao_duplica_nem_regrava(tmp_path):
    xl = criar_planilha(tmp_path / "g.xlsx")
    reg = [{"sg": "333333", "valor_total": Decimal("10.00")}]
    Planilha(cfg(xl, tmp_path)).upsert(reg)
    antes = sha(xl)
    acoes = Planilha(cfg(xl, tmp_path)).upsert(reg)
    assert [a.acao for a in acoes] == ["sem_mudanca"]
    assert sha(xl) == antes
    ws = load_workbook(xl)["Garantia"]
    assert [ws.cell(row=r, column=1).value for r in range(4, 8)] == [111111, 222222, 333333, None]


def test_sg_existente_preenche_so_vazio(tmp_path):
    xl = criar_planilha(tmp_path / "g.xlsx")
    acoes = Planilha(cfg(xl, tmp_path)).upsert([{"sg": "000222222", "valor_total": Decimal("50")}])
    assert acoes[0].acao == "atualizada" and acoes[0].linha == 5
    ws = load_workbook(xl)["Garantia"]
    assert ws["G5"].value == 50.0
    assert ws.max_row == 5 or ws["A6"].value is None


def test_sg_existente_com_valor_diferente_e_mantida(tmp_path):
    xl = criar_planilha(tmp_path / "g.xlsx")
    antes = sha(xl)
    acoes = Planilha(cfg(xl, tmp_path)).upsert([{"sg": "111111", "valor_total": Decimal("999")}])
    assert acoes[0].acao == "divergente" and "999" in acoes[0].detalhe
    assert sha(xl) == antes


def test_ao_divergir_atualizar(tmp_path):
    xl = criar_planilha(tmp_path / "g.xlsx")
    Planilha(cfg(xl, tmp_path, ao_divergir="atualizar")).upsert([{"sg": "111111", "valor_total": Decimal("999")}])
    assert load_workbook(xl)["Garantia"]["G4"].value == 999.0


def test_valor_igual_em_formato_diferente_nao_e_divergencia(tmp_path):
    xl = criar_planilha(tmp_path / "g.xlsx")
    acoes = Planilha(cfg(xl, tmp_path)).upsert([{"sg": "111111", "valor_total": Decimal("300.00")}])
    assert acoes[0].acao == "sem_mudanca"


def test_backup_uma_vez_por_sessao(tmp_path):
    xl = criar_planilha(tmp_path / "g.xlsx")
    original = sha(xl)
    p = Planilha(cfg(xl, tmp_path))
    p.upsert([{"sg": "1", "valor_total": Decimal("1")}])
    p.upsert([{"sg": "2", "valor_total": Decimal("2")}])
    backups = list((tmp_path / "bkp").glob("*.xlsx"))
    assert len(backups) == 1 and sha(backups[0]) == original


def test_tabela_do_excel_cresce(tmp_path):
    xl = criar_planilha(tmp_path / "g.xlsx", com_tabela=True)
    Planilha(cfg(xl, tmp_path)).upsert([{"sg": "333333", "valor_total": Decimal("1")}])
    assert load_workbook(xl)["Garantia"].tables["Garantia"].ref == "A3:H6"


def test_planilha_aberta_no_excel(tmp_path):
    xl = criar_planilha(tmp_path / "g.xlsx")
    (tmp_path / "~$g.xlsx").write_text("trava")
    with pytest.raises(ErroPlanilha, match="aberta no Excel"):
        Planilha(cfg(xl, tmp_path)).upsert([{"sg": "1", "valor_total": Decimal("1")}])


def test_nao_cria_planilha_nova(tmp_path):
    with pytest.raises(ErroPlanilha, match="nao encontrada"):
        Planilha(cfg(tmp_path / "nao_existe.xlsx", tmp_path)).upsert([{"sg": "1"}])
    assert not (tmp_path / "nao_existe.xlsx").exists()


def test_recusa_planilha_com_grafico(tmp_path):
    from openpyxl.chart import BarChart, Reference

    xl = criar_planilha(tmp_path / "g.xlsx")
    wb = load_workbook(xl)
    ws = wb["Garantia"]
    graf = BarChart()
    graf.add_data(Reference(ws, min_col=7, min_row=3, max_row=5), titles_from_data=True)
    ws.add_chart(graf, "J2")
    wb.save(xl)
    antes = sha(xl)
    with pytest.raises(ErroPlanilha, match="graficos"):
        Planilha(cfg(xl, tmp_path)).upsert([{"sg": "1", "valor_total": Decimal("1")}])
    assert sha(xl) == antes


def test_sem_coluna_chave(tmp_path):
    xl = tmp_path / "vazia.xlsx"
    Workbook().save(xl)
    with pytest.raises(ErroPlanilha, match="cabecalho"):
        Planilha(cfg(xl, tmp_path)).upsert([{"sg": "1"}])


def test_caminho_nao_configurado(tmp_path):
    with pytest.raises(ErroPlanilha, match="SAGA_PLANILHA"):
        ConfigPlanilha.from_dict({"caminho": "", "colunas": COLUNAS}, TIPOS, tmp_path)


def test_linha_de_total_nao_e_sobrescrita_nem_vira_modelo(tmp_path, caplog):
    xl = criar_planilha(tmp_path / "g.xlsx")
    wb = load_workbook(xl)
    ws = wb["Garantia"]
    ws["F6"], ws["G6"] = "TOTAL", "=SUM(G4:G5)"
    wb.save(xl)
    Planilha(cfg(xl, tmp_path)).upsert([{"sg": "333333", "valor_total": Decimal("5")}])
    ws = load_workbook(xl)["Garantia"]
    assert ws["G6"].value == "=SUM(G4:G5)" and ws["F6"].value == "TOTAL"
    assert ws["A7"].value == 333333 and ws["G7"].value == 5.0
    assert ws["H7"].value == "=E7+F7-G7"  # formula veio da linha 5 (com SG), nao da de total
    assert "total?" in caplog.text
