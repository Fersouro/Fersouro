"""Camada de relatorios: definicao em YAML -> pasta de trabalho .xlsx."""

from __future__ import annotations

import datetime as dt
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from openpyxl import load_workbook

from datalake.config import ConfigError
from datalake.report import (
    FORMATS,
    ReportConfig,
    build_all,
    column_formats,
    infer_format,
    load_reports,
)

RELATORIO_YML = """
name: pedidos
title: Pedidos por cliente
description: Relatorio de teste.
formats:
  vlr_total: moeda
sheets:
  - name: Resumo
    description: Total por cliente.
    sql: |
      SELECT nome, sum(vlr_total) AS vlr_total, count(*) AS qtd_pedidos
        FROM pedidos_cliente
       GROUP BY nome
       ORDER BY nome
    totals: [vlr_total, qtd_pedidos]
    highlights:
      - when: vlr_total > 200
        style: amarelo
  - name: Detalhe
    model: pedidos_cliente
"""


@pytest.fixture
def lake_com_gold(project, settings, control):
    """Roda ingest -> silver -> gold para ter um modelo materializado."""
    from datalake.layers import bronze, gold, silver

    bronze.ingest_source(settings, settings.source("erp"), control, "t1")
    silver.build_source(settings, settings.source("erp"), control, "t1")
    gold.build_all(settings, control, "t1")
    return settings


def _escrever_relatorio(project, conteudo=RELATORIO_YML, nome="pedidos.yml"):
    pasta = project / "conf" / "reports"
    pasta.mkdir(parents=True, exist_ok=True)
    (pasta / nome).write_text(conteudo, encoding="utf-8")


# ------------------------------------------------------------------ formatos


def test_dinheiro_pelo_nome_da_coluna():
    """O tipo nao distingue: float chamado venda_total e dinheiro."""
    assert infer_format("venda_total", [Decimal("10.5")]) == FORMATS["moeda"]
    assert infer_format("preco_medio", [1.5]) == FORMATS["moeda"]


def test_revenda_e_numero_da_loja_nao_dinheiro():
    """'revenda' contem 'venda' -- nao pode virar R$."""
    assert infer_format("revenda", [1, 2]) == FORMATS["inteiro"]


def test_porcentagem_nao_multiplica_de_novo():
    """A gold ja entrega 12.5 para 12,5%; o formato '%' do Excel multiplicaria."""
    formato = infer_format("lucro_porcentagem", [12.5])
    assert formato == FORMATS["percentual"]
    assert "%" in formato and not formato.endswith("0.0%")


def test_data_e_competencia():
    assert infer_format("emissao", [dt.date(2026, 8, 13)]) == FORMATS["data"]
    assert infer_format("competencia", [dt.datetime(2026, 8, 1)]) == FORMATS["mes"]
    assert infer_format("dta_documento", [dt.datetime(2026, 8, 13, 9, 30)]) == FORMATS["data_hora"]


def test_texto_e_nulo_ficam_sem_formato():
    assert infer_format("descricao", ["parafuso"]) is None
    assert infer_format("qualquer", [None, None]) is None


def test_formato_declarado_vence_a_inferencia():
    formatos = column_formats(["valor"], [(Decimal("1"),)], {"valor": "inteiro"})
    assert formatos == [FORMATS["inteiro"]]


def test_mascara_crua_do_excel_passa_direto():
    assert column_formats(["x"], [(1.0,)], {"x": '#,##0.000'}) == ['#,##0.000']


# ----------------------------------------------------------------- definicao


def test_model_vira_select_estrela():
    cfg = ReportConfig.from_dict(
        {"name": "r", "sheets": [{"name": "A", "model": "margem_pecas"}]}
    )
    assert cfg.sheets[0].sql == 'SELECT * FROM "margem_pecas"'


def test_sql_e_model_juntos_e_erro():
    with pytest.raises(ConfigError, match="nao os dois"):
        ReportConfig.from_dict(
            {"name": "r", "sheets": [{"name": "A", "sql": "SELECT 1", "model": "x"}]}
        )


def test_aba_sem_origem_e_erro():
    with pytest.raises(ConfigError, match="falta 'sql' ou 'model'"):
        ReportConfig.from_dict({"name": "r", "sheets": [{"name": "A"}]})


def test_estilo_invalido_e_erro():
    with pytest.raises(ConfigError, match="estilo"):
        ReportConfig.from_dict(
            {
                "name": "r",
                "sheets": [
                    {"name": "A", "sql": "SELECT 1", "highlights": [{"when": "1=1", "style": "roxo"}]}
                ],
            }
        )


def test_relatorio_sem_abas_e_erro():
    with pytest.raises(ConfigError, match="sem 'sheets'"):
        ReportConfig.from_dict({"name": "r"})


def test_load_reports_le_a_pasta(project, settings):
    _escrever_relatorio(project)
    relatorios = load_reports(settings)
    assert [r.name for r in relatorios] == ["pedidos"]
    assert relatorios[0].title == "Pedidos por cliente"


def test_sem_pasta_de_relatorios_nao_falha(settings):
    assert load_reports(settings) == []
    assert build_all(settings) == []


def test_relatorio_inexistente(project, settings):
    _escrever_relatorio(project)
    with pytest.raises(ValueError, match="inexistente"):
        build_all(settings, apenas=["nao_existe"])


# -------------------------------------------------------------------- saida


def test_gera_xlsx_com_capa_totais_e_destaque(project, lake_com_gold, tmp_path):
    _escrever_relatorio(project)
    destino = tmp_path / "saida"
    resultados = build_all(lake_com_gold, destino_dir=destino)

    assert [r.status for r in resultados] == ["success"]
    caminho = resultados[0].path
    assert caminho == destino / "pedidos.xlsx"

    wb = load_workbook(caminho)
    # A capa vem primeiro: quem recebe a planilha sabe o que e e de quando.
    assert wb.sheetnames == ["Capa", "Resumo", "Detalhe"]
    assert wb["Capa"]["A1"].value == "Pedidos por cliente"

    ws = wb["Resumo"]
    assert [c.value for c in ws[1]] == ["nome", "vlr_total", "qtd_pedidos"]
    assert ws.freeze_panes == "A2"
    assert ws["B2"].number_format == FORMATS["moeda"]
    assert ws["C2"].number_format == FORMATS["inteiro"]

    # Uma linha de total, com SUBTOTAL para acompanhar o filtro.
    ultima = ws.max_row
    assert ws.cell(ultima, 1).value == "TOTAL"
    assert ws.cell(ultima, 2).value.startswith("=SUBTOTAL(109,")
    assert ws.auto_filter.ref.endswith(str(ultima - 1))   # o total fica fora do filtro


def test_destaque_pinta_a_linha_que_bate_na_condicao(project, lake_com_gold, tmp_path):
    _escrever_relatorio(project)
    wb = load_workbook(build_all(lake_com_gold, destino_dir=tmp_path)[0].path)
    ws = wb["Resumo"]

    pintadas = {
        ws.cell(linha, 1).value
        for linha in range(2, ws.max_row)
        if ws.cell(linha, 1).fill.fgColor.rgb == "00FFEB9C"
    }
    # 'Beta' tem pedido de 250,50 -- unico acima de 200.
    assert pintadas == {"Beta"}


def test_destaque_em_uma_coluna_so(project, lake_com_gold, tmp_path):
    _escrever_relatorio(
        project,
        """
name: escopo
sheets:
  - name: Resumo
    sql: SELECT nome, vlr_total FROM pedidos_cliente ORDER BY nome
    highlights:
      - when: vlr_total > 200
        style: vermelho
        scope: vlr_total
""",
        "escopo.yml",
    )
    resultado = [r for r in build_all(lake_com_gold, destino_dir=tmp_path) if r.report == "escopo"][0]
    ws = load_workbook(resultado.path)["Resumo"]
    linha = next(i for i in range(2, ws.max_row + 1) if ws.cell(i, 1).value == "Beta")
    assert ws.cell(linha, 2).fill.fgColor.rgb == "00FFC7CE"
    assert ws.cell(linha, 1).fill.fgColor.rgb != "00FFC7CE"     # so a coluna pedida


def test_modelo_ausente_e_skipped_nao_falha(project, settings, tmp_path):
    """Relatorio que cita modelo ainda nao carregado nao derruba a execucao."""
    _escrever_relatorio(
        project,
        "name: futuro\nsheets:\n  - name: A\n    model: nao_carregado\n",
        "futuro.yml",
    )
    resultados = build_all(settings, destino_dir=tmp_path)
    assert [r.status for r in resultados] == ["skipped"]
    assert all(r.ok for r in resultados)


def test_total_de_coluna_inexistente_falha_com_mensagem(project, lake_com_gold, tmp_path):
    _escrever_relatorio(
        project,
        "name: ruim\nsheets:\n  - name: A\n    model: pedidos_cliente\n    totals: [nao_existe]\n",
        "ruim.yml",
    )
    resultado = [r for r in build_all(lake_com_gold, destino_dir=tmp_path) if r.report == "ruim"][0]
    assert resultado.status == "failed"
    assert "nao_existe" in resultado.message


def test_corte_no_teto_do_excel_conta_o_que_ficou_de_fora(lake_com_gold):
    """A aba grava o que cabe e diz o total real -- sem carregar tudo."""
    from datalake.duck import connect
    from datalake.layers.gold import register_silver_views
    from datalake.report import SheetConfig, fetch_sheet, register_gold_views

    con = connect(lake_com_gold)
    register_silver_views(con, lake_com_gold)
    register_gold_views(con, lake_com_gold)
    try:
        sheet = SheetConfig(name="A", sql="SELECT * FROM pedidos_cliente")
        colunas, linhas, _, total = fetch_sheet(con, sheet, max_rows=1)
        assert len(linhas) == 1
        assert total == 2            # dois pedidos nao cancelados na base de teste
    finally:
        con.close()


def test_limit_declarado_nao_dispara_contagem(lake_com_gold):
    from datalake.duck import connect
    from datalake.layers.gold import register_silver_views
    from datalake.report import SheetConfig, fetch_sheet, register_gold_views

    con = connect(lake_com_gold)
    register_silver_views(con, lake_com_gold)
    register_gold_views(con, lake_com_gold)
    try:
        sheet = SheetConfig(name="A", sql="SELECT * FROM pedidos_cliente", limit=1)
        _, linhas, _, total = fetch_sheet(con, sheet, max_rows=1000)
        assert len(linhas) == 1 and total == 1     # limite pedido nao e aviso
    finally:
        con.close()


# ---------------------------------------------------------------- parametros


def _cfg_com_parametros(**extra):
    base = {
        "name": "r",
        "parameters": [
            {"name": "competencia", "label": "Competencia", "type": "mes", "default": "atual"},
            {"name": "departamento", "type": "numero", "default": 410, "optional": True},
        ],
        "sheets": [{"name": "A", "sql": "SELECT 1"}],
    }
    base.update(extra)
    return ReportConfig.from_dict(base)


def test_competencia_atual_vira_primeiro_dia_do_mes():
    hoje = dt.date.today()
    valores = _cfg_com_parametros().resolve_parameters()
    assert valores["competencia"] == dt.date(hoje.year, hoje.month, 1)


def test_competencia_aceita_os_formatos_que_a_pessoa_digita():
    cfg = _cfg_com_parametros()
    esperado = dt.date(2026, 3, 1)
    for texto in ("2026-03", "03/2026", "2026-03-17"):
        assert cfg.resolve_parameters({"competencia": texto})["competencia"] == esperado


def test_competencia_invalida_diz_o_formato():
    with pytest.raises(ValueError, match="AAAA-MM"):
        _cfg_com_parametros().resolve_parameters({"competencia": "agosto"})


def test_numero_invalido_e_recusado():
    with pytest.raises(ValueError, match="nao e um numero"):
        _cfg_com_parametros().resolve_parameters({"departamento": "funilaria"})


def test_opcional_em_branco_vira_nulo():
    """Departamento vazio no formulario = todos os departamentos."""
    assert _cfg_com_parametros().resolve_parameters({"departamento": ""})["departamento"] is None


def test_obrigatorio_sem_valor_e_sem_default_falha():
    cfg = ReportConfig.from_dict({
        "name": "r",
        "parameters": [{"name": "filial", "type": "texto"}],
        "sheets": [{"name": "A", "sql": "SELECT 1"}],
    })
    with pytest.raises(ValueError, match="obrigatorio"):
        cfg.resolve_parameters({})


def test_parametro_que_o_relatorio_nao_tem():
    with pytest.raises(ValueError, match="inexistente"):
        _cfg_com_parametros().resolve_parameters({"filial": "1"})


def test_tipo_de_parametro_invalido_no_yaml():
    with pytest.raises(ConfigError, match="tipo 'cor' invalido"):
        ReportConfig.from_dict({
            "name": "r",
            "parameters": [{"name": "x", "type": "cor"}],
            "sheets": [{"name": "A", "sql": "SELECT 1"}],
        })


def test_parametro_filtra_de_verdade_e_aparece_na_capa(project, lake_com_gold, tmp_path):
    """O valor vai como parametro do DuckDB ($nome), nao concatenado no SQL."""
    _escrever_relatorio(
        project,
        """
name: por_cliente
parameters:
  - name: cliente
    label: Cliente
    type: texto
    default: Alfa
sheets:
  - name: Resumo
    sql: SELECT nome, vlr_total FROM pedidos_cliente WHERE nome = $cliente
""",
        "por_cliente.yml",
    )
    resultados = {r.report: r for r in build_all(lake_com_gold, ["por_cliente"],
                                                 tmp_path, {"cliente": "Beta"})}
    resultado = resultados["por_cliente"]
    assert resultado.status == "success"

    wb = load_workbook(resultado.path)
    nomes = [wb["Resumo"].cell(l, 1).value for l in range(2, wb["Resumo"].max_row + 1)]
    assert nomes == ["Beta"]                       # filtrou pelo valor pedido

    capa = [c.value for linha in wb["Capa"].iter_rows(values_only=False) for c in linha]
    assert "Cliente" in capa and "Beta" in capa     # a planilha diz o que foi filtrado


def test_atalhos_de_data_no_default():
    """'inicio-do-mes' e 'fim-do-mes' poupam digitar a data toda geracao."""
    cfg = ReportConfig.from_dict({
        "name": "r",
        "parameters": [
            {"name": "de", "type": "data", "default": "inicio-do-mes"},
            {"name": "ate", "type": "data", "default": "fim-do-mes"},
            {"name": "quando", "type": "data", "default": "hoje"},
        ],
        "sheets": [{"name": "A", "sql": "SELECT 1"}],
    })
    valores = cfg.resolve_parameters()
    hoje = dt.date.today()
    assert valores["de"] == hoje.replace(day=1)
    assert valores["quando"] == hoje
    assert valores["ate"].month == hoje.month
    assert (valores["ate"] + dt.timedelta(days=1)).month != hoje.month   # ultimo dia


def test_fim_do_mes_em_dezembro(monkeypatch):
    """Dezembro e o caso que quebra quem soma um mes sem pensar."""
    import datalake.report as report

    class Dezembro(dt.date):
        @classmethod
        def today(cls):
            return dt.date(2026, 12, 7)

    monkeypatch.setattr(report.dt, "date", Dezembro)
    assert report._fim_do_mes() == dt.date(2026, 12, 31)


def test_data_aceita_os_dois_formatos_e_recusa_o_resto():
    cfg = ReportConfig.from_dict({
        "name": "r",
        "parameters": [{"name": "de", "type": "data", "default": "hoje"}],
        "sheets": [{"name": "A", "sql": "SELECT 1"}],
    })
    assert cfg.resolve_parameters({"de": "2026-09-30"})["de"] == dt.date(2026, 9, 30)
    assert cfg.resolve_parameters({"de": "30/09/2026"})["de"] == dt.date(2026, 9, 30)
    with pytest.raises(ValueError, match="AAAA-MM-DD"):
        cfg.resolve_parameters({"de": "setembro"})


def test_periodo_filtra_pelas_datas_escolhidas(project, lake_com_gold, tmp_path):
    _escrever_relatorio(
        project,
        """
name: periodo
parameters:
  - name: data_inicial
    label: Data inicial
    type: data
    default: inicio-do-mes
  - name: data_final
    label: Data final
    type: data
    default: fim-do-mes
sheets:
  - name: Pedidos
    sql: |
      SELECT id_pedido, vlr_total FROM pedidos_cliente
       WHERE CAST(id_pedido AS INTEGER) BETWEEN 10 AND 11
         AND $data_inicial <= $data_final
""",
        "periodo.yml",
    )
    resultado = [r for r in build_all(lake_com_gold, ["periodo"], tmp_path,
                                      {"data_inicial": "2026-01-01", "data_final": "2026-01-31"})][0]
    assert resultado.status == "success"
    capa = [c.value for linha in load_workbook(resultado.path)["Capa"].iter_rows() for c in linha]
    assert "01/01/2026" in capa and "31/01/2026" in capa


def test_historico_de_estoque_vira_view(project, settings, tmp_path):
    """O ERP so tem o saldo de agora; o historico diario e a unica serie que existe."""
    import duckdb

    from datalake.duck import connect
    from datalake.report import register_history_views

    pasta = settings.root / "historico_estoque"
    pasta.mkdir(parents=True)
    escritor = duckdb.connect()
    escritor.execute(
        "CREATE TABLE s(data DATE, revenda INTEGER, codigo VARCHAR, disponivel DOUBLE)"
    )
    escritor.execute("INSERT INTO s VALUES ('2026-09-08', 1, 'X', 5)")
    escritor.execute(f"COPY s TO '{pasta / '2026-09-08.parquet'}' (FORMAT PARQUET)")
    escritor.close()

    con = connect(settings)
    try:
        assert register_history_views(con, settings) == ["historico_estoque"]
        assert con.execute("SELECT disponivel FROM historico_estoque").fetchone()[0] == 5
    finally:
        con.close()


def test_sem_historico_nao_registra_nada(settings):
    from datalake.duck import connect
    from datalake.report import register_history_views

    con = connect(settings)
    try:
        assert register_history_views(con, settings) == []
    finally:
        con.close()


# ------------------------------------------------------- parametro de escolha


def _cfg_revenda():
    return ReportConfig.from_dict({
        "name": "r",
        "parameters": [{
            "name": "revenda", "label": "Revenda", "type": "lista",
            "default": "", "optional": True,
            "options": [{"value": 1, "label": "Revenda 1"},
                        {"value": 2, "label": "Revenda 2"},
                        {"value": "", "label": "Consolidado (1 e 2)"}],
        }],
        "sheets": [{"name": "A", "sql": "SELECT 1"}],
    })


def test_escolha_vira_numero_para_o_sql():
    assert _cfg_revenda().resolve_parameters({"revenda": "2"})["revenda"] == 2


def test_consolidado_e_ausencia_de_filtro():
    """Consolidado nao e um terceiro codigo de loja: e nao filtrar loja nenhuma."""
    assert _cfg_revenda().resolve_parameters({"revenda": ""})["revenda"] is None
    assert _cfg_revenda().resolve_parameters({})["revenda"] is None      # default


def test_opcao_fora_da_lista_e_recusada():
    with pytest.raises(ValueError, match="nao e uma opcao valida"):
        _cfg_revenda().resolve_parameters({"revenda": "7"})


def test_lista_sem_options_e_erro_de_configuracao():
    with pytest.raises(ConfigError, match="exige 'options'"):
        ReportConfig.from_dict({
            "name": "r",
            "parameters": [{"name": "x", "type": "lista"}],
            "sheets": [{"name": "A", "sql": "SELECT 1"}],
        })


def test_capa_mostra_o_rotulo_escolhido(project, lake_com_gold, tmp_path):
    """Na capa vale 'Consolidado', nao o valor tecnico que foi para o SQL."""
    _escrever_relatorio(
        project,
        """
name: escolha
parameters:
  - name: cliente
    label: Cliente
    type: lista
    default: ""
    optional: true
    options:
      - value: Alfa
        label: Só Alfa
      - value: ""
        label: Todos os clientes
sheets:
  - name: Resumo
    sql: |
      SELECT nome FROM pedidos_cliente
       WHERE ($cliente IS NULL OR nome = $cliente)
       ORDER BY nome
""",
        "escolha.yml",
    )
    resultado = [r for r in build_all(lake_com_gold, ["escolha"], tmp_path)][0]
    capa = [c.value for linha in load_workbook(resultado.path)["Capa"].iter_rows()
            for c in linha]
    assert "Todos os clientes" in capa
